"""Generate DOCX/ODT cases; requires python-docx, Pillow and Node.js."""

import io
import zipfile

from common import encrypt, finish, options
from common import save as save_file
from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw

args = options(__doc__, node=True)


def save(name, data):
    save_file(args.output, name, data)


def word(empty=False, image=False):
    doc = Document()
    if not empty and (not image):
        doc.add_heading("Synthetic coursework report", 0)
        doc.add_paragraph(
            "This document checks file handling only. "
            "It is not a completed coursework report."
        )
    if image:
        picture = Image.new("RGB", (600, 180), "white")
        ImageDraw.Draw(picture).text(
            (20, 60), "Synthetic scanned report page", fill="black"
        )
        data = io.BytesIO()
        picture.save(data, format="PNG")
        data.seek(0)
        doc.add_picture(data, width=Inches(4))
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def change(data, remove=None, replacements=None):
    output = io.BytesIO()
    replacements = replacements or {}
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(
        output, "w", zipfile.ZIP_DEFLATED
    ) as out:
        for info in original.infolist():
            if info.filename != remove:
                out.writestr(info, replacements.get(info.filename, original.read(info)))
    return output.getvalue()


valid = word()
picture = word(image=True)
save("14-text-report.docx", valid)
save("15-empty-report.docx", word(empty=True))
save("16-image-only-report.docx", picture)
save("17-truncated-report.docx", valid[:-12])
save("18-missing-main-part.docx", change(valid, remove="word/document.xml"))
save("19-missing-image-part.docx", change(picture, remove="word/media/image1.png"))
save(
    "20-malformed-document.docx",
    change(valid, replacements={"word/document.xml": b"<w:document>"}),
)
save("21-text-renamed-docx.docx", b"This is plain text, not a Word document.")
save("22-docx-renamed-doc.doc", valid)


def odt(empty=False):
    body = "" if empty else "<text:p>Synthetic report with written text.</text:p>"
    content = (
        '<?xml version="1.0" encoding="UTF-8"?><office:document-content '
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
        'office:version="1.2"><office:body><office:text>'
        f"{body}</office:text></office:body></office:document-content>"
    )
    manifest = (
        '<?xml version="1.0"?><manifest:manifest '
        'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        'manifest:version="1.2"><manifest:file-entry manifest:full-path="/" '
        'manifest:media-type="application/vnd.oasis.opendocument.text"/>'
        '<manifest:file-entry manifest:full-path="content.xml" '
        'manifest:media-type="text/xml"/></manifest:manifest>'
    )
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as out:
        out.writestr(
            "mimetype",
            "application/vnd.oasis.opendocument.text",
            compress_type=zipfile.ZIP_STORED,
        )
        out.writestr("content.xml", content, compress_type=zipfile.ZIP_DEFLATED)
        out.writestr(
            "META-INF/manifest.xml", manifest, compress_type=zipfile.ZIP_DEFLATED
        )
    return output.getvalue()


valid_odt = odt()
save("23-text-report.odt", valid_odt)
save("24-empty-report.odt", odt(empty=True))
save("26-odt-renamed-docx.docx", valid_odt)
save("27-missing-content.odt", change(valid_odt, remove="content.xml"))
encrypt(args, valid, "25-password-protected-report.docx")
finish(args, "Reports")
